#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从 data/考点库.json 生成 Excel 清单：输出/上海高中物理-考点公式清单.xlsx"""
import json
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
DATA = ROOT / "data" / "考点库.json"
OUT = ROOT / "输出" / "上海高中物理-考点公式清单.xlsx"

GREEN = "0B6B5B"
HEAD_FILL = PatternFill("solid", fgColor=GREEN)
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
ZEBRA = PatternFill("solid", fgColor="F4F8F6")
THIN = Side(style="thin", color="D8D8D8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_sheet(ws, widths, wrap_cols=()):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=(c.column in wrap_cols))
            if row_idx % 2 == 0:
                c.fill = ZEBRA
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    d = json.loads(DATA.read_text(encoding="utf-8"))
    units = d["units"]
    OUT.parent.mkdir(exist_ok=True)

    wb = Workbook()

    # --- 1 考点总表 ---
    ws = wb.active
    ws.title = "考点总表"
    ws.append(["编号", "考点", "所属单元", "要求", "考频", "重点", "例题数", "先猜一句", "源文件"])
    for u in units:
        for c in u["cards"]:
            ws.append([c["id"], c["name"], c["unit"], c["req"], c["freq"], c["star"],
                       len(c["examples"]), c["guess"], c["file"]])
    style_sheet(ws, [10, 34, 30, 6, 6, 6, 8, 46, 34], wrap_cols=(2, 3, 8, 9))

    # --- 2 例题索引 ---
    ws = wb.create_sheet("例题索引")
    ws.append(["考点", "例题", "难度", "标题", "题干（截断）"])
    for u in units:
        for c in u["cards"]:
            for e in c["examples"]:
                ws.append([c["id"], f"例{e['n']}", e["level"], e["title"], e["problem"]])
    style_sheet(ws, [10, 8, 8, 24, 76], wrap_cols=(4, 5))

    # --- 3 公式表 ---
    ws = wb.create_sheet("公式表")
    ws.append(["单元", "#", "公式", "含义", "适用条件"])
    for u in units:
        for f in u["formulas"]:
            ws.append([f["unit"], f["no"], f["formula"], f["meaning"], f["cond"]])
    style_sheet(ws, [30, 6, 42, 34, 30], wrap_cols=(1, 3, 4, 5))

    # --- 4 关系边 ---
    ws = wb.create_sheet("关系边")
    ws.append(["起点", "终点", "关系类型", "单元"])
    for u in units:
        for r in u["relations"]:
            ws.append([r["from"], r["to"], r["type"], u["title"]])
    style_sheet(ws, [12, 26, 16, 32], wrap_cols=(2, 4))

    # --- 5 关卡地图 ---
    ws = wb.create_sheet("关卡地图")
    ws.append(["单元", "关", "考点", "难度", "星级", "状态"])
    for u in units:
        for lv in u["levels"]:
            ws.append([u["title"], lv["no"], lv["name"], lv["diff"], lv["star"], "☐"])
    style_sheet(ws, [32, 6, 30, 8, 10, 8], wrap_cols=(1, 3))

    # --- 6 说明 ---
    ws = wb.create_sheet("说明")
    s = d["stats"]
    rows = [
        ["《上海高中物理知识图谱》考点与公式清单"],
        [""],
        ["统计", f"{s['units']} 个单元 / {s['cards']} 个考点 / {s['examples']} 道例题 / {s['formulas']} 条公式"],
        ["来源", "由 build/parse.py 从 Markdown 知识卡自动解析，源文件变更后重跑即可"],
        ["用法", "考点总表按单元筛选；例题索引可当题库；公式表可打印成速查卡；关卡地图可做打卡"],
        ["", ""],
        ["版本", "v5（自评 + 分层 + 趣味层 + 限时 + 错题复仇）"],
    ]
    for r in rows:
        ws.append(r)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 90
    ws["A1"].font = Font(size=15, bold=True, color=GREEN)

    wb.save(OUT)
    print("saved", OUT, OUT.stat().st_size, "bytes")


if __name__ == "__main__":
    main()

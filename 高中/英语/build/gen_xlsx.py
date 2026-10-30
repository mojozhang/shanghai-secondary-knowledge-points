# -*- coding: utf-8 -*-
"""生成 Excel：单元总表 / 必背词 / 词汇总表 / 故事。"""
import json
import pathlib
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
CARDS = json.loads((ROOT / "data" / "考点库.json").read_text(encoding="utf-8"))
VOC = json.loads((ROOT / "data" / "词汇库.json").read_text(encoding="utf-8"))
STORIES = json.loads((ROOT / "data" / "故事.json").read_text(encoding="utf-8"))
OUT = ROOT / "输出" / "上海高中英语-考点词汇清单.xlsx"
GREEN = "0B6B5B"


def head(ws, cols, widths):
    ws.append(cols)
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor=GREEN)
        c.font = Font(color="FFFFFF", bold=True, size=10.5)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "单元总表"
    head(ws, ["编号", "单元", "册", "语法", "写作", "必背词数"],
         [8, 26, 14, 34, 40, 10])
    best = {}
    for b in CARDS["books"]:
        for c in b["cards"]:
            m = re.search(r"本单元必背\s*\d+\s*词\s*\n+\s*(.+)", c["body"])
            words = [w.strip() for w in m.group(1).split(",") if w.strip()] if m else []
            best[c["id"]] = words
            ws.append([c["id"], c["name"], b["title"], "", "", len(words)])

    ws = wb.create_sheet("必背词")
    head(ws, ["编号", "单元", "单词", "词性", "中文", "D0", "D+1", "D+3", "D+7", "D+15"],
         [8, 22, 22, 10, 46, 6, 6, 6, 6, 6])
    index = {}
    for r in VOC:
        index.setdefault(r["word"].lower(), r)
    for cid, words in best.items():
        for w in words:
            r = index.get(w.lower(), {})
            ws.append([cid, r.get("id", ""), w, r.get("pos_all", ""), r.get("cn", ""),
                       "", "", "", "", ""])

    ws = wb.create_sheet("词汇总表")
    head(ws, ["#", "册", "单元", "英文", "词性", "中文", "标星", "D0", "D+1", "D+3", "D+7", "D+15"],
         [5, 8, 12, 22, 10, 60, 6, 6, 6, 6, 6, 6])
    for i, r in enumerate(VOC, 1):
        ws.append([i, r["book_no"], r["unit"], r["word"], r["pos_all"], r["cn"],
                   "★" if r["star"] else "", "", "", "", "", ""])

    ws = wb.create_sheet("故事")
    head(ws, ["册", "标题", "章数", "英文词数", "中文梗概"], [8, 29, 6, 10, 70])
    for k, s in sorted(STORIES.items(), key=lambda x: int(x[0])):
        words = sum(len(ch["text"].replace("**", "").split()) for ch in s["chapters"])
        ws.append([int(k), s["title"], len(s["chapters"]), words, s.get("cn", "")])

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print("saved", OUT, OUT.stat().st_size, "bytes")


if __name__ == "__main__":
    main()

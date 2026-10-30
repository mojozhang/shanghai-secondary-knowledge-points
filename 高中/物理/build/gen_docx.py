#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从 data/考点库.json 生成 Word 手册：输出/上海高中物理知识图谱手册.docx

排版：封面 → 目录 → 各单元（考点卡片：标题、元信息、概述、自评、原理、例题、配图、表格）
"""
import json
import pathlib
import re
import subprocess

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import mathfmt

ROOT = pathlib.Path(__file__).resolve().parent.parent
DATA = ROOT / "data" / "考点库.json"
OUTDIR = ROOT / "输出"
DOCX = OUTDIR / "上海高中物理知识图谱手册.docx"
PNG_DIR = ROOT / "图" / "png"

SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
GREEN = RGBColor(0x0B, 0x6B, 0x5B)
GRAY = RGBColor(0x6B, 0x6B, 0x70)


def set_font(run, size=9.5, bold=False, color=None, ea="宋体", ascii_="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = ascii_
    run._element.rPr.rFonts.set(qn("w:eastAsia"), ea)
    if color is not None:
        run.font.color.rgb = color


def para(doc, text="", size=9.5, bold=False, color=None, align=None,
         space_before=0.5, space_after=0.5, indent=None, ea="宋体"):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.12
    if indent:
        pf.left_indent = Cm(indent)
    add_inline(p, text, size=size, bold=bold, color=color, ea=ea)
    return p


EMOJI_RE = re.compile("[\u26a1\u23f1\U0001f4a5\U0001f511\U0001f48e\U0001f9ea\U0001f3c6\U0001f525\U0001f4ca\U0001f5fa\U0001f432\U0001f3ae\U0001f4cd\U0001f512\u2705\u2b50]")


def add_inline(p, text, size=9.5, bold=False, color=None, ea="宋体"):
    text = EMOJI_RE.sub("", text)
    for seg in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            add_rich(p, seg[2:-2], size, True, color, ea=ea)
        elif seg.startswith("`") and seg.endswith("`"):
            r = p.add_run(seg[1:-1])
            set_font(r, size, bold, color, ea="Consolas", ascii_="Consolas")
        else:
            add_rich(p, seg, size, bold, color, ea=ea)


def add_rich(p, text, size=9.5, bold=False, color=None, ea="宋体"):
    """普通文字与公式混排：公式片段渲染为内嵌图片。"""
    pos = 0
    for a, b, frag in mathfmt.spans(text):
        if a > pos:
            r = p.add_run(text[pos:a])
            set_font(r, size, bold, color, ea=ea)
        r = mathfmt.render(frag)
        if r:
            png, _, h_cm = r
            run = p.add_run()
            run.add_picture(str(png), height=Cm(0.42))
        else:
            r2 = p.add_run(frag)
            set_font(r2, size, bold, color, ea=ea)
        pos = b
    if pos < len(text):
        r = p.add_run(text[pos:])
        set_font(r, size, bold, color, ea=ea)


def add_table(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncol):
            txt = row[j] if j < len(row) else ""
            txt = EMOJI_RE.sub("", re.sub(r"\*\*([^*]+)\*\*", r"\1", txt).replace("`", ""))
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if mathfmt.is_formula(txt):
                rr = mathfmt.render(txt)
                if rr:
                    p.add_run().add_picture(str(rr[0]), height=Cm(0.38))
                else:
                    r = p.add_run(txt)
                    set_font(r, 8, bold=(i == 0), ea="宋体")
            else:
                add_rich(p, txt, 8, (i == 0), ea="宋体")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_image(doc, rel_path, width_cm=9.5):
    direct = ROOT / rel_path
    if str(rel_path).lower().endswith(".png") and direct.exists():
        png = direct
    else:
        name = pathlib.Path(rel_path).stem + ".png"
        png = PNG_DIR / name
    if not png.exists():
        para(doc, f"[缺图：{rel_path}]", 9, color=GRAY)
        return
    # 竖图按高度约束，横图按宽度约束，避免一张图吃掉半页
    try:
        from PIL import Image as _Img
        w, h = _Img.open(png).size
        if h > w:
            width_cm = min(width_cm, 6.2)
    except Exception:
        pass
    doc.add_picture(str(png), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def convert_png():
    svgs = list((ROOT / "图").glob("*.svg"))
    if not svgs:
        return
    PNG_DIR.mkdir(parents=True, exist_ok=True)
    todo = [s for s in svgs if not (PNG_DIR / (s.stem + ".png")).exists()]
    if todo:
        subprocess.run([SOFFICE, "--headless", "--convert-to", "png",
                        "--outdir", str(PNG_DIR)] + [str(s) for s in todo],
                       check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def render_body(doc, body):
    """把卡片正文（Markdown 片段）渲染进 docx。"""
    lines = body.splitlines()
    i, buf = 0, []

    def flush_para():
        nonlocal buf
        text = " ".join(x.strip() for x in buf).strip()
        buf = []
        if text:
            para(doc, text, 9.5, space_before=1, space_after=1)

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()

        if s.startswith("```"):
            flush_para()
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                r = doc.add_paragraph()
                r.paragraph_format.space_before = Pt(0)
                r.paragraph_format.space_after = Pt(0)
                r.paragraph_format.left_indent = Cm(0.6)
                run = r.add_run(lines[i])
                set_font(run, 9, ea="Consolas", ascii_="Consolas")
                i += 1
            i += 1
            continue

        if s.startswith("|"):
            flush_para()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") and c for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue

        m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", s)
        if m:
            flush_para()
            add_image(doc, m.group(1))
            i += 1
            continue

        if s.startswith("## "):
            flush_para()
            p = para(doc, "", space_before=9, space_after=3)
            add_inline(p, s[3:], 12, True, RGBColor(0x0A, 0x4F, 0x43), ea="微软雅黑")
            i += 1
            continue

        if s.startswith("### "):
            flush_para()
            p = para(doc, "", space_before=6, space_after=2)
            add_inline(p, s[4:], 11, True, GREEN, ea="微软雅黑")
            i += 1
            continue

        if s.startswith(">"):
            flush_para()
            text = re.sub(r"^>\s?", "", s)
            if text.strip():
                para(doc, text, 8.5, color=GRAY, indent=0.5, space_before=0.5, space_after=0.5)
            i += 1
            continue

        if re.match(r"^[-*] ", s):
            flush_para()
            para(doc, "· " + s[2:], 9.5, indent=0.4, space_before=0, space_after=0)
            i += 1
            continue

        if re.match(r"^\d+\. ", s):
            flush_para()
            para(doc, s, 9.5, indent=0.4, space_before=0, space_after=0)
            i += 1
            continue

        if s.startswith("*") and s.endswith("*") and len(s) > 2 and s.count("*") == 2:
            flush_para()
            para(doc, s.strip("*"), 8, color=GRAY, space_before=0, space_after=3)
            i += 1
            continue

        if s in ("---", "***", "___"):
            flush_para()
            i += 1
            continue

        if not s:
            flush_para()
            i += 1
            continue

        buf.append(ln)
        i += 1
    flush_para()


def main():
    d = json.loads(DATA.read_text(encoding="utf-8"))
    units = d["units"]
    convert_png()
    OUTDIR.mkdir(exist_ok=True)

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.4)
    sec.bottom_margin = Cm(1.4)
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)

    # 封面
    for _ in range(5):
        doc.add_paragraph()
    p = para(doc, "上海高中物理", 24, True, GREEN, WD_ALIGN_PARAGRAPH.CENTER, ea="微软雅黑")
    p = para(doc, "知 识 图 谱", 24, True, GREEN, WD_ALIGN_PARAGRAPH.CENTER, ea="微软雅黑")
    doc.add_paragraph()
    para(doc, "按沪科版（2020）教材体系编排 · 适合高一至高三", 12, False, GRAY,
         WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "自评 · 分层 · 例题 · 限时 · 错题回炉", 12, False, GRAY, WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    s = d["stats"]
    para(doc, f"{s['units']} 个单元　{s['cards']} 个考点　{s['examples']} 道例题　{s['formulas']} 条公式",
         11, False, GRAY, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "2026-10-30", 10, False, GRAY, WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # 怎么用这本书
    guide = ROOT / "00a-怎么用这本书.md"
    if guide.exists():
        gt = guide.read_text(encoding="utf-8")
        gt = re.sub(r"^# .*$", "", gt, count=1, flags=re.M)
        para(doc, "怎么用这本书", 20, True, GREEN, WD_ALIGN_PARAGRAPH.CENTER, ea="微软雅黑")
        doc.add_paragraph()
        render_body(doc, gt)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # 目录
    p = para(doc, "目　录", 15, True, GREEN, WD_ALIGN_PARAGRAPH.CENTER, ea="微软雅黑")
    doc.add_paragraph()
    for u in units:
        para(doc, u["title"], 10, True, ea="微软雅黑", space_before=4)
        for c in u["cards"]:
            para(doc, f"{c['id']}　{c['name']}", 8.5, indent=0.5, space_before=0, space_after=0)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # 正文
    for ui, u in enumerate(units):
        if ui:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        add_inline(p, u["title"], 15, True, GREEN, ea="微软雅黑")

        for c in u["cards"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            add_inline(p, f"{c['id']}  {c['name']}", 12.5, True, RGBColor(0x0A, 0x4F, 0x43), ea="微软雅黑")
            if c["meta"]:
                para(doc, c["meta"], 7.5, color=GRAY, space_before=0, space_after=4)

            body = c["body"]
            # 去掉元信息行（已在上面单独渲染）
            blines = body.split("\n")
            for k, ln in enumerate(blines):
                if ln.strip():
                    if ln.strip().startswith("*") and ln.strip().endswith("*"):
                        blines[k] = ""
                    break
            body = "\n".join(blines)
            render_body(doc, body)

    doc.save(DOCX)
    print("saved", DOCX, DOCX.stat().st_size, "bytes")

    import os
    if os.environ.get("DSH_SKIP_PDF") == "1":
        return
    pdf = OUTDIR / "上海高中物理知识图谱手册.pdf"
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf",
                    "--outdir", str(OUTDIR), str(DOCX)],
                   check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print("pdf:", pdf.exists(), pdf.stat().st_size if pdf.exists() else 0)


if __name__ == "__main__":
    main()
